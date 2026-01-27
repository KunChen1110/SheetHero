"""Excel chart helpers."""

import io
import tempfile
from typing import Optional

import matplotlib.pyplot as plt
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, ScatterChart
from openpyxl.chart.reference import Reference
from openpyxl.drawing.image import Image
from PIL import Image as PILImage


class ExcelChartManager:
    """Chart creation and embedding helpers."""

    def __init__(self, workbook, reader, temp_files: list):
        self.workbook = workbook
        self.reader = reader
        self._temp_files = temp_files

    def save_plot_to_excel(self, sheet_name: str, cell_position: str = "A1",
                           figsize: tuple = (10, 6), dpi: int = 100) -> str:
        """
        Save current matplotlib plot as image in Excel sheet.
        """
        if sheet_name not in self.workbook.sheetnames:
            self.workbook.create_sheet(sheet_name)
        sheet = self.workbook[sheet_name]

        fig = plt.gcf()
        if fig.get_axes():
            fig.set_size_inches(figsize)
            plt.tight_layout()

            img_buffer = io.BytesIO()
            plt.savefig(img_buffer, format='png', dpi=dpi, bbox_inches='tight')
            img_buffer.seek(0)

            pil_img = PILImage.open(img_buffer)

            with tempfile.NamedTemporaryFile(suffix='.png', delete=False) as tmp_file:
                pil_img.save(tmp_file.name, 'PNG')
                tmp_filename = tmp_file.name

            img = Image(tmp_filename)
            sheet.add_image(img, cell_position)

            self._temp_files.append(tmp_filename)

            print(
                f"✅ Chart saved to sheet '{sheet_name}' at position {cell_position}"
            )
            plt.close(fig)
            return f"Chart saved to {sheet_name}!{cell_position}"

        print("⚠️ No plot found to save. Create a plot first.")
        return "No plot to save"

    def create_chart(self, sheet_name: str, chart_type: str, data_range: str,
                     position: str = "A1", title: str = "",
                     x_axis_title: str = "", y_axis_title: str = "") -> str:
        """
        Create chart in Excel sheet.
        """
        try:
            sheet = self.reader.get_sheet(sheet_name)

            chart_classes = {
                'bar': BarChart,
                'line': LineChart,
                'pie': PieChart,
                'scatter': ScatterChart,
                'area': AreaChart
            }

            if chart_type.lower() not in chart_classes:
                raise ValueError(
                    f"Unsupported chart type: {chart_type}. Available: {list(chart_classes.keys())}"
                )

            chart_class = chart_classes[chart_type.lower()]
            chart = chart_class()

            if title:
                chart.title = title
            if x_axis_title and hasattr(chart, 'x_axis'):
                chart.x_axis.title = x_axis_title
            if y_axis_title and hasattr(chart, 'y_axis'):
                chart.y_axis.title = y_axis_title

            data = Reference(sheet, range_string=data_range)
            chart.add_data(data, titles_from_data=True)

            sheet.add_chart(chart, position)

            message = (
                f"✅ Created {chart_type} chart from {data_range} at {position} in sheet '{sheet_name}'"
            )
            print(message)
            return message

        except Exception as e:
            error_msg = f"❌ Error creating chart: {str(e)}"
            print(error_msg)
            raise Exception(error_msg)
